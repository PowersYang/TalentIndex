import os
import pandas as pd
from docx import Document
from win32com import client as wc


def save_as_docx(path):
    "doc文档转docx"
    count = 0
    word = wc.Dispatch('Word.Application')
    try:
        for dirpath, dirnames, files in os.walk(path):
            if files:
                for filename in files:
                    if filename.endswith('.doc') or filename.endswith('.docx'):
                        file_path = dirpath + '\\' + filename
                        new_file_path = 'F:\\Mine\\PycharmProjects' + '\\docx\\' + filename.split('.')[0] + '.docx'
                        print(file_path)
                        print(new_file_path)

                        doc = word.Documents.Open(file_path)  # 目标路径下的文件
                        doc.SaveAs(new_file_path, 12, False, "", True, "", False, False, False,
                                   False)  # 转化后路径下的文件
                        doc.SaveAs()
                        doc.Close()
                        count += 1
    except:
        print(file_path)

    word.Quit()

    print(count)


def get_data(path):
    columns = ["企业名称", "当年月平均职工总数", "签订劳动合同关系且大专以上学历职工人数", "占当年月平均职工总数比例（%）", "签订劳动合同关系且大专以上学历的研发人员数",
               "占当年月平均职工总数比例（%）", "研发人员总数", "管理人员总数", "市场推广人员总数", "大专及本科学历职工数", "硕士学历职工数", "博士学历职工数"]

    lis = []
    for dirpath, dirnames, files in os.walk(path):
        for filename in files:
            file_path = dirpath + '\\' + filename
            doc = Document(file_path)
            try:
                table = doc.tables[5]
            except IndexError:
                print(filename)

            try:
                a = table.cell(1, 1).text
            except:
                print('a: ' + filename)
                a = 0

            try:
                b = table.cell(1, 3).text
            except:
                print('b: ' + filename)
                b = 0

            try:
                c = table.cell(1, 5).text
            except:
                print('c: ' + filename)
                c = 0

            try:
                d = table.cell(1, 7).text
            except:
                print('d: ' + filename)
                d = 0

            try:
                e = table.cell(1, 9).text
            except:
                print('e: ' + filename)
                e = 0

            try:
                f = table.cell(3, 1).text
            except:
                print('f: ' + filename)
                f = 0

            try:
                g = table.cell(3, 3).text
            except:
                print('g: ' + filename)
                g = 0

            try:
                h = table.cell(3, 5).text
            except:
                print('h: ' + filename)
                h = 0

            try:
                i = table.cell(3, 7).text
            except:
                print('i: ' + filename)
                i = 0

            try:
                j = table.cell(3, 9).text
            except:
                print('j: ' + filename)
                j = 0

            try:
                k = table.cell(3, 10).text
            except:
                print('k: ' + filename)
                k = 0

            temp_data = [filename[:-5], a, b, c, d, e, f, g, h, i, j, k]
            lis.append(temp_data)

    df = pd.DataFrame(lis, columns=columns)
    df.to_excel('talentIndex.xls')


if __name__ == '__main__':
    get_data("F:\Mine\PycharmProjects\docx")
